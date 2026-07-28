#!/usr/bin/env python3
"""Build the Chinese RTL RS-check architecture design DOCX.

Requires python-docx and Pillow. The Markdown source remains authoritative;
this script applies the repository's Word layout, tables, code blocks, and
figure pagination deterministically.
"""

from __future__ import annotations

import argparse
import re
import tempfile
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor, Twips


INK = RGBColor(32, 45, 58)
NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(92, 104, 116)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CODE_FILL = "F6F8FA"
WHITE = "FFFFFF"

PAGE_WIDTH = Inches(8.5)
PAGE_HEIGHT = Inches(11)
LANDSCAPE_WIDTH = Inches(11)
LANDSCAPE_HEIGHT = Inches(8.5)
MARGIN = Inches(1)
HEADER_FOOTER_DISTANCE = Inches(0.492)
PORTRAIT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}

FIGURE_CONFIG = {
    "rs-tool-overall-architecture.png": {"orientation": "landscape", "width": 8.7},
    "rs-tool-core-data-model.png": {"orientation": "landscape", "width": 7.4},
    "rs-tool-end-to-end-workflow.png": {"orientation": "portrait", "width": 5.25, "parts": 2},
    "rs-tool-online-collection-sequence.png": {"orientation": "landscape", "width": 8.7},
    "rs-tool-crg-trace-workflow.png": {"orientation": "portrait", "width": 5.7},
    "rs-tool-row-check-workflow.png": {"orientation": "portrait", "width": 6.05, "parts": 3},
    "rs-tool-timeout-cancellation-sequence.png": {"orientation": "landscape", "width": 6.8},
}


@dataclass(frozen=True)
class Element:
    kind: str
    value: object
    level: int = 0


def _is_table_separator(line: str) -> bool:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def _table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def parse_markdown(path: Path) -> list[Element]:
    lines = path.read_text(encoding="utf-8").splitlines()
    elements: list[Element] = []
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue

        if line.startswith("```"):
            language = line[3:].strip()
            index += 1
            block: list[str] = []
            while index < len(lines) and not lines[index].startswith("```"):
                block.append(lines[index])
                index += 1
            if index >= len(lines):
                raise ValueError(f"unterminated fenced code block in {path}")
            elements.append(Element("code", {"language": language, "text": "\n".join(block)}))
            index += 1
            continue

        image = re.fullmatch(r"!\[([^]]+)]\(([^)]+)\)", line.strip())
        if image:
            elements.append(Element("image", {"alt": image.group(1), "path": image.group(2)}))
            index += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            elements.append(Element("heading", heading.group(2).strip(), len(heading.group(1))))
            index += 1
            continue

        if line.lstrip().startswith("|") and index + 1 < len(lines) and _is_table_separator(lines[index + 1]):
            rows = [_table_row(line)]
            index += 2
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                rows.append(_table_row(lines[index]))
                index += 1
            width = len(rows[0])
            if any(len(row) != width for row in rows):
                raise ValueError(f"inconsistent Markdown table width near row {index}")
            elements.append(Element("table", rows))
            continue

        ordered = re.match(r"^\s*\d+\.\s+(.+)$", line)
        bullet = re.match(r"^\s*-\s+(.+)$", line)
        if ordered or bullet:
            list_kind = "ordered_list" if ordered else "bullet_list"
            marker_pattern = r"^\s*\d+\.\s+(.+)$" if ordered else r"^\s*-\s+(.+)$"
            items: list[str] = []
            while index < len(lines):
                match = re.match(marker_pattern, lines[index])
                if not match:
                    break
                item = match.group(1).strip()
                index += 1
                continuations: list[str] = []
                while (
                    index < len(lines)
                    and lines[index].strip()
                    and not re.match(r"^\s*(?:\d+\.|-)\s+", lines[index])
                    and not lines[index].startswith("#")
                    and not lines[index].startswith("```")
                    and not lines[index].lstrip().startswith("|")
                ):
                    continuations.append(lines[index].strip())
                    index += 1
                if continuations:
                    item += " " + " ".join(continuations)
                items.append(item)
                while index < len(lines) and not lines[index].strip():
                    index += 1
            elements.append(Element(list_kind, items))
            continue

        paragraph_lines = [line.strip()]
        index += 1
        while index < len(lines) and lines[index].strip():
            candidate = lines[index]
            if (
                candidate.startswith("#")
                or candidate.startswith("```")
                or re.fullmatch(r"!\[([^]]+)]\(([^)]+)\)", candidate.strip())
                or re.match(r"^\s*(?:\d+\.|-)\s+", candidate)
                or (candidate.lstrip().startswith("|") and index + 1 < len(lines) and _is_table_separator(lines[index + 1]))
            ):
                break
            paragraph_lines.append(candidate.strip())
            index += 1
        elements.append(Element("paragraph", " ".join(paragraph_lines)))
    return elements


def set_run_font(
    run,
    *,
    ascii_name: str = "Calibri",
    east_asia: str = "Microsoft YaHei",
    size: float | None = None,
    color: RGBColor | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = ascii_name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), ascii_name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), ascii_name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    lang = run._element.get_or_add_rPr().find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run._element.get_or_add_rPr().append(lang)
    lang.set(qn("w:eastAsia"), "zh-CN")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_run_shading(run, fill: str) -> None:
    r_pr = run._element.get_or_add_rPr()
    shd = r_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        r_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _ensure_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def apply_table_geometry(table, widths_dxa: Sequence[int]) -> None:
    widths = [int(value) for value in widths_dxa]
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = _ensure_child(tbl_pr, "w:tblW")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total))
    tbl_ind = _ensure_child(tbl_pr, "w:tblInd")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    layout = _ensure_child(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for col_index, width in enumerate(widths):
        table.columns[col_index].width = Twips(width)
    for row in table.rows:
        row.height = None
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for col_index, cell in enumerate(row.cells):
            cell.width = Twips(widths[col_index])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = _ensure_child(tc_pr, "w:tcW")
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[col_index]))
            tc_mar = _ensure_child(tc_pr, "w:tcMar")
            for side, margin in CELL_MARGINS_DXA.items():
                node = _ensure_child(tc_mar, f"w:{side}")
                node.set(qn("w:type"), "dxa")
                node.set(qn("w:w"), str(margin))


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    heading_specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    if "Figure Caption" not in doc.styles:
        caption = doc.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = doc.styles["Figure Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.font.size = Pt(9.5)
    caption.font.color.rgb = MUTED
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_together = True

    if "Code Block" not in doc.styles:
        code = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = doc.styles["Code Block"]
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei UI")
    code.font.size = Pt(8.5)
    code.font.color.rgb = INK
    code.paragraph_format.left_indent = Inches(0.16)
    code.paragraph_format.right_indent = Inches(0.08)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.keep_together = True


def set_section_layout(section, landscape: bool) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = LANDSCAPE_WIDTH
        section.page_height = LANDSCAPE_HEIGHT
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = PAGE_WIDTH
        section.page_height = PAGE_HEIGHT
    section.top_margin = MARGIN
    section.right_margin = MARGIN
    section.bottom_margin = MARGIN
    section.left_margin = MARGIN
    section.header_distance = HEADER_FOOTER_DISTANCE
    section.footer_distance = HEADER_FOOTER_DISTANCE


def add_field(run, instruction: str, placeholder: str = "1") -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)


def configure_header_footer(section) -> None:
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("RTL 打拍例化检查工具设计方案")
    set_run_font(run, size=8.5, color=MUTED)

    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=MUTED)
    page = paragraph.add_run()
    set_run_font(page, size=8.5, color=MUTED)
    add_field(page, "PAGE")
    run = paragraph.add_run(" 页，共 ")
    set_run_font(run, size=8.5, color=MUTED)
    total = paragraph.add_run()
    set_run_font(total, size=8.5, color=MUTED)
    add_field(total, "NUMPAGES")
    run = paragraph.add_run(" 页")
    set_run_font(run, size=8.5, color=MUTED)


def configure_all_section_links(doc: Document) -> None:
    for section in doc.sections[1:]:
        section.header.is_linked_to_previous = True
        section.footer.is_linked_to_previous = True


def add_hyperlink(paragraph, text: str, target: str) -> None:
    relationship = paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend((color, underline))
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend((r_pr, text_node))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(r"(`[^`]+`|\[[^]]+]\([^)]+\)|\*\*[^*]+\*\*)")


def add_inline(paragraph, text: str, *, size: float = 11, bold: bool = False) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, size=size, color=INK, bold=bold)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(
                run,
                ascii_name="Consolas",
                east_asia="Microsoft YaHei UI",
                size=max(8.5, size - 0.8),
                color=DARK_BLUE,
                bold=False,
            )
            set_run_shading(run, LIGHT_GRAY)
        elif token.startswith("["):
            link = re.fullmatch(r"\[([^]]+)]\(([^)]+)\)", token)
            if link:
                add_hyperlink(paragraph, link.group(1), link.group(2))
        else:
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=INK, bold=True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=size, color=INK, bold=bold)


def _next_abstract_id(numbering) -> int:
    ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    return max(ids, default=-1) + 1


def _next_num_id(numbering) -> int:
    ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    return max(ids, default=0) + 1


def create_abstract_numbering(doc: Document, ordered: bool) -> int:
    numbering = doc.part.numbering_part.element
    abstract_id = _next_abstract_id(numbering)
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1." if ordered else "•")
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    p_pr.extend((tabs, indent))
    level.extend((start, num_fmt, lvl_text, suff, p_pr))
    abstract.append(level)
    numbering.append(abstract)
    return abstract_id


def create_num_id(doc: Document, abstract_id: int) -> int:
    numbering = doc.part.numbering_part.element
    num_id = _next_num_id(numbering)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_id))
    num.append(abstract)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend((ilvl, num_id_node))


def add_list(doc: Document, items: Iterable[str], ordered: bool, abstract_id: int) -> None:
    num_id = create_num_id(doc, abstract_id)
    for item in items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.25
        paragraph.paragraph_format.widow_control = True
        apply_num(paragraph, num_id)
        add_inline(paragraph, item)


def table_widths(header: Sequence[str]) -> list[int]:
    key = tuple(cell.strip() for cell in header)
    if key == ("内部字段", "行内要求", "用途"):
        return [1900, 1800, 5660]
    if key == ("类型", "示例", "是否使检查失败"):
        return [1900, 5000, 2460]
    if key == ("模块", "主要职责"):
        return [3450, 5910]
    if key == ("名称", "设计含义"):
        return [2450, 6910]
    if key == ("根节点", "内容"):
        return [2750, 6610]
    if key == ("拍数参数状态", "拍数贡献"):
        return [6700, 2660]
    if len(header) == 2:
        return [2900, 6460]
    widths = [PORTRAIT_WIDTH_DXA // len(header)] * len(header)
    widths[-1] += PORTRAIT_WIDTH_DXA - sum(widths)
    return widths


def add_table(doc: Document, rows: Sequence[Sequence[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    apply_table_geometry(table, table_widths(rows[0]))
    mark_header_row(table.rows[0])
    for row_index, source_row in enumerate(rows):
        for col_index, value in enumerate(source_row):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            if row_index == 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_cell_shading(cell, LIGHT_BLUE)
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                set_cell_shading(cell, WHITE)
            add_inline(paragraph, value, size=9.5, bold=row_index == 0)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_code_block(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Code Block")
    set_paragraph_shading(paragraph, CODE_FILL)
    p_pr = paragraph._p.get_or_add_pPr()
    borders = _ensure_child(p_pr, "w:pBdr")
    left = _ensure_child(borders, "w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "10")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), "2E74B5")
    run = paragraph.add_run(text)
    set_run_font(
        run,
        ascii_name="Consolas",
        east_asia="Microsoft YaHei UI",
        size=8.5,
        color=INK,
    )


def add_title_page(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(28)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(8)
    run = kicker.add_run("技术设计文档")
    set_run_font(run, size=11, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("RTL 打拍例化检查工具设计方案")
    set_run_font(run, size=25, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(24)
    run = subtitle.add_run("架构、RTL 清单采集、检查算法与图形界面设计")
    set_run_font(run, size=13, color=MUTED)

    rows = [
        ("项目", "内容"),
        ("文档类型", "工具设计方案"),
        ("适用范围", "Excel 规格解析、Verdi KDB 采集、RTL 打拍例化检查"),
        ("结构版本", "RTL 清单 v3；JSON 报告 v4"),
        ("运行环境", "Python 3.8+；在线采集使用 Linux 与 Verdi/NPI"),
        ("更新日期", date.today().isoformat()),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    apply_table_geometry(table, [1850, 7510])
    mark_header_row(table.rows[0])
    for row_index, (label, value) in enumerate(rows):
        for col_index, text in enumerate((label, value)):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, LIGHT_BLUE if row_index == 0 or col_index == 0 else WHITE)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(paragraph, text, size=10, bold=row_index == 0 or col_index == 0)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(20)
    note.paragraph_format.space_after = Pt(0)
    note.paragraph_format.line_spacing = 1.2
    set_paragraph_shading(note, LIGHT_GRAY)
    run = note.add_run(
        "说明：本文描述当前已实现并签核的工具行为。产品名、命令、配置键、结构字段、"
        "类名和检查项代码保留原始拼写，以便与源码、日志和输出文件直接对照。"
    )
    set_run_font(run, size=10, color=INK)
    doc.add_page_break()


def add_toc(doc: Document) -> None:
    heading = doc.add_paragraph(style="Heading 1")
    add_inline(heading, "目录", size=16, bold=True)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    set_run_font(run, size=10.5, color=INK)
    add_field(run, 'TOC \\o "1-3" \\h \\z \\u', "右键更新目录")
    doc.add_page_break()


def find_split_boundaries(image: Image.Image, parts: int) -> list[int]:
    gray = image.convert("L")
    width, height = gray.size
    boundaries: list[int] = []
    for part in range(1, parts):
        target = height * part // parts
        low = max(boundaries[-1] + 400 if boundaries else 400, target - 350)
        high = min(height - 400, target + 350)
        best_y = target
        best_ink: int | None = None
        for y in range(low, high + 1, 3):
            row = gray.crop((20, y, width - 20, y + 1))
            ink = sum(row.histogram()[:245])
            if best_ink is None or ink < best_ink:
                best_ink = ink
                best_y = y
        boundaries.append(best_y)
    return boundaries


def split_image(path: Path, parts: int, directory: Path) -> list[Path]:
    with Image.open(path) as source:
        image = source.convert("RGB")
        boundaries = find_split_boundaries(image, parts)
        edges = [0, *boundaries, image.height]
        outputs: list[Path] = []
        overlap = 0
        for index in range(parts):
            top = max(0, edges[index] - (overlap if index else 0))
            bottom = min(image.height, edges[index + 1] + (overlap if index + 1 < parts else 0))
            target = directory / f"{path.stem}-part-{index + 1}.png"
            image.crop((0, top, image.width, bottom)).save(target, optimize=True)
            outputs.append(target)
        return outputs


def set_image_alt(inline_shape, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", description)
    doc_pr.set("descr", description)


def add_figure_part(doc: Document, path: Path, width_in: float, caption: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    shape = run.add_picture(str(path), width=Inches(width_in))
    set_image_alt(shape, caption)
    cap = doc.add_paragraph(style="Figure Caption")
    add_inline(cap, caption, size=9.5)


def add_new_section(doc: Document, landscape: bool) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_layout(section, landscape)
    section.header.is_linked_to_previous = True
    section.footer.is_linked_to_previous = True


def add_figure(
    doc: Document,
    source_dir: Path,
    image_data: dict[str, str],
    figure_number: int,
    temporary_dir: Path,
) -> None:
    path = (source_dir / image_data["path"]).resolve()
    if not path.is_file():
        raise FileNotFoundError(f"figure not found: {path}")
    config = FIGURE_CONFIG.get(path.name)
    if config is None:
        raise ValueError(f"missing figure layout configuration for {path.name}")
    parts = int(config.get("parts", 1))
    part_paths = split_image(path, parts, temporary_dir) if parts > 1 else [path]
    for index, part_path in enumerate(part_paths, start=1):
        if index > 1:
            doc.add_page_break()
        suffix = f"（第 {index}/{parts} 部分）" if parts > 1 else ""
        caption = f"图 {figure_number} {image_data['alt']}{suffix}"
        add_figure_part(doc, part_path, float(config["width"]), caption)


def add_body_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.widow_control = True
    add_inline(paragraph, text)


def render_elements(doc: Document, elements: Sequence[Element], source: Path, temp_dir: Path) -> None:
    bullet_abstract = create_abstract_numbering(doc, ordered=False)
    ordered_abstract = create_abstract_numbering(doc, ordered=True)
    figure_number = 0
    index = 0
    while index < len(elements):
        element = elements[index]
        if element.kind == "heading" and element.level == 1:
            index += 1
            continue

        next_is_image = index + 1 < len(elements) and elements[index + 1].kind == "image"
        if element.kind == "heading" and next_is_image:
            image_data = elements[index + 1].value
            image_path = (source.parent / image_data["path"]).resolve()
            config = FIGURE_CONFIG.get(image_path.name)
            if config is None:
                raise ValueError(f"missing figure configuration for {image_path.name}")
            landscape = config["orientation"] == "landscape"
            if landscape:
                add_new_section(doc, landscape=True)
            else:
                doc.add_page_break()
            style = "Heading 1" if element.level == 2 else "Heading 2"
            paragraph = doc.add_paragraph(style=style)
            add_inline(paragraph, str(element.value), size=16 if style == "Heading 1" else 13, bold=True)
            figure_number += 1
            add_figure(doc, source.parent, image_data, figure_number, temp_dir)
            if landscape:
                add_new_section(doc, landscape=False)
            else:
                doc.add_page_break()
            index += 2
            continue

        if element.kind == "heading":
            style_level = max(1, min(3, element.level - 1))
            style = f"Heading {style_level}"
            paragraph = doc.add_paragraph(style=style)
            size = {1: 16, 2: 13, 3: 12}[style_level]
            add_inline(paragraph, str(element.value), size=size, bold=True)
        elif element.kind == "paragraph":
            add_body_paragraph(doc, str(element.value))
        elif element.kind == "bullet_list":
            add_list(doc, element.value, ordered=False, abstract_id=bullet_abstract)
        elif element.kind == "ordered_list":
            add_list(doc, element.value, ordered=True, abstract_id=ordered_abstract)
        elif element.kind == "table":
            add_table(doc, element.value)
        elif element.kind == "code":
            add_code_block(doc, element.value["text"])
        elif element.kind == "image":
            figure_number += 1
            add_figure(doc, source.parent, element.value, figure_number, temp_dir)
        else:
            raise ValueError(f"unsupported Markdown element: {element.kind}")
        index += 1


def enable_field_updates(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def build(source: Path, output: Path) -> None:
    elements = parse_markdown(source)
    doc = Document()
    configure_styles(doc)
    set_section_layout(doc.sections[0], landscape=False)
    configure_header_footer(doc.sections[0])
    doc.core_properties.title = "RTL 打拍例化检查工具设计方案"
    doc.core_properties.subject = "RTL 打拍例化检查工具架构与检查流程"
    doc.core_properties.author = "RTL Verification Team"
    doc.core_properties.keywords = "RTL, NPI, Verdi, Excel, 打拍检查"
    doc.core_properties.comments = "由仓库中文设计文档生成"
    enable_field_updates(doc)

    add_title_page(doc)
    add_toc(doc)
    with tempfile.TemporaryDirectory(prefix="rs-tool-docx-figures-") as name:
        render_elements(doc, elements, source, Path(name))
    configure_all_section_links(doc)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "docs" / "ARCHITECTURE_AND_DESIGN.md",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "docs" / "RTL_RS_CHECK_TOOL_DESIGN_CN.docx",
    )
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
